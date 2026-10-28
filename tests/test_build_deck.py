"""Tests for assembling the deck.

The point of the manifest is that this script never learns what a metric is, so the
tests lean on that: a made-up chart type and a made-up switch have to work without the
builder recognising either.
"""

import os
import tempfile

import yaml
from PIL import Image
from pptx import Presentation

import _bootstrap
from build_deck import (
    GRID_CELLS,
    PER_SLIDE,
    SINGLE_CELL,
    STACKED_CELLS,
    build_deck,
    cells_for,
    group_by_metric,
    group_by_scope,
    paginate,
    _fit,
    _heading_for,
)

PICTURE = 13  # MSO_SHAPE_TYPE.PICTURE


def chart(name, scope, title=None, metric=None):
    entry = {"file": name, "title": title or name, "scope": scope}
    if metric is not None:
        entry["metric"] = metric
    return entry


def write_manifest(folder, charts, test_type="soak", sizes=None):
    for entry in charts:
        size = (sizes or {}).get(entry["file"], (400, 300))
        Image.new("RGB", size, "white").save(os.path.join(folder, entry["file"]))
    manifest = {"capture": "x.yml", "test_type": test_type,
                "ber_source": "fec_counter", "charts": charts}
    path = os.path.join(folder, "manifest.yml")
    with open(path, "w", encoding="utf-8") as f:
        yaml.dump(manifest, f, sort_keys=False, allow_unicode=True)
    return path


def slide_pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == PICTURE]


# --- grouping follows the manifest, it does not sort ------------------------------

groups = group_by_scope([chart("a.png", "all"), chart("b.png", "all"),
                         chart("c.png", "0"), chart("d.png", "4-33")])
assert [scope for scope, _ in groups] == ["all", "0", "4-33"], groups
assert len(groups[0][1]) == 2
print("ok  group_by_scope: manifest order kept, consecutive scopes collected")

# the same scope appearing again later stays a separate group rather than being merged,
# because reordering would contradict the order the runner chose
groups = group_by_scope([chart("a.png", "0"), chart("b.png", "all"), chart("c.png", "0")])
assert [scope for scope, _ in groups] == ["0", "all", "0"]
print("ok  group_by_scope: order wins over tidiness")


# --- pagination ------------------------------------------------------------------

assert paginate([1, 2, 3, 4, 5], 4) == [[1, 2, 3, 4], [5]]
assert paginate([], 4) == []
assert paginate([1], 4) == [[1]]
print("ok  paginate: splits on the grid size, no empty trailing page")

assert _heading_for("all", None, 1, 1) == "Across all switches"
assert _heading_for("0", None, 1, 1) == "Switch 0"
assert _heading_for("0", "Pre-FEC BER", 1, 1) == "Switch 0 — Pre-FEC BER"
assert _heading_for("0", "Pre-FEC BER", 2, 3) == "Switch 0 — Pre-FEC BER (page 2 of 3)"
print("ok  _heading_for: names the metric, pages only when there is more than one")


# --- a slide holds one metric, and its charts share the room ---------------------

groups = group_by_metric([chart("a.png", "0", metric="BER"),
                          chart("b.png", "0", metric="BER"),
                          chart("c.png", "0", metric="T-Code")])
assert [m for m, _ in groups] == ["BER", "T-Code"], groups
assert len(groups[0][1]) == 2
print("ok  group_by_metric: consecutive charts of one metric share a slide")

groups = group_by_metric([chart("a.png", "0"), chart("b.png", "0")])
assert len(groups) == 1 and groups[0][0] is None,     "a manifest with no metric field must stay one group, as it laid out before"
print("ok  group_by_metric: a manifest without the field is unchanged")

assert cells_for(1) is SINGLE_CELL
assert cells_for(2) is STACKED_CELLS
assert cells_for(3) is GRID_CELLS and cells_for(4) is GRID_CELLS
assert STACKED_CELLS[0][3] > GRID_CELLS[0][3], "two stacked charts keep the full width"
print("ok  cells_for: room per chart follows how many are sharing the slide")


# --- pictures are fitted, never stretched ----------------------------------------

work = tempfile.mkdtemp()
wide = os.path.join(work, "wide.png")
Image.new("RGB", (1200, 300), "white").save(wide)
width, height = _fit(wide, 5.9, 2.25)
assert abs(width / height - 4.0) < 0.01, "a 4:1 image must stay 4:1"
assert width <= 5.9 + 1e-6 and height <= 2.25 + 1e-6, "and must fit inside the cell"

tall = os.path.join(work, "tall.png")
Image.new("RGB", (300, 1200), "white").save(tall)
width, height = _fit(tall, 5.9, 2.25)
assert height <= 2.25 + 1e-6, "a tall image is limited by height, not width"
assert abs(width / height - 0.25) < 0.01
print("ok  _fit: aspect preserved, bounded by whichever side runs out first")


# --- one chart takes the slide, several share the grid ---------------------------

work = tempfile.mkdtemp()
path = write_manifest(work, [chart("solo.png", "0")])
build_deck(path, os.path.join(work, "one.pptx"))
deck = Presentation(os.path.join(work, "one.pptx"))
only = slide_pictures(deck.slides[2])[0]
assert only.width > GRID_CELLS[0][3] * 914400, \
    "a lone chart should fill the slide, not sit in a quarter of it"
assert SINGLE_CELL[0][3] > GRID_CELLS[0][3]
print("ok  layout: a single chart gets the full slide")

work = tempfile.mkdtemp()
path = write_manifest(work, [chart(f"c{i}.png", "all") for i in range(PER_SLIDE)])
build_deck(path, os.path.join(work, "grid.pptx"))
deck = Presentation(os.path.join(work, "grid.pptx"))
assert len(slide_pictures(deck.slides[2])) == PER_SLIDE
print(f"ok  layout: {PER_SLIDE} charts share one slide")


# --- the builder must not need to recognise anything -----------------------------

work = tempfile.mkdtemp()
charts = [chart("ber.png", "all"),
          chart("brand_new_metric.png", "all"),      # a type added upstream later
          chart("temp_switch_99-7.png", "99-7")]     # a switch named anything
path = write_manifest(work, charts, test_type="shut_unshut")
build_deck(path, os.path.join(work, "unknown.pptx"))
deck = Presentation(os.path.join(work, "unknown.pptx"))
headings = []
for slide in deck.slides:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            headings.append(shape.text_frame.text.strip())
            break
assert "Shut_Unshut Plots" in headings, headings
assert "Switch 99-7" in headings, "an unrecognised switch id still gets its own slide"
total = sum(len(slide_pictures(s)) for s in deck.slides)
assert total == 3, "every chart in the manifest is placed, known or not"
print("ok  unknown metrics, switches and test types all survive untouched")


# --- an empty capture produces a short deck, not a crash -------------------------

work = tempfile.mkdtemp()
path = write_manifest(work, [])
build_deck(path, os.path.join(work, "empty.pptx"))
deck = Presentation(os.path.join(work, "empty.pptx"))
assert len(deck.slides) == 2, "title and section divider only"
print("ok  a capture with no charts still builds")

print()
print("all tests passed")
