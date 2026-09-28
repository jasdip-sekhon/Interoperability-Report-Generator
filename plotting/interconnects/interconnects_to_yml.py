from pathlib import Path
from openpyxl import load_workbook
import yaml

def interconnects_to_yml(interconnects_file, output_yml):
    wb = load_workbook(interconnects_file)
    worksheet = wb.active

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return {}

    headers = rows[0]
    data_rows = rows[1:]

    interconnects = []
    for row in data_rows:
        if not any(row):
            continue
        record = {}
        for i in range(len(headers)): 
            header = headers[i]
            if i < len(row):
                record[header] = row[i]
        interconnects.append(record)

    with open(output_yml, 'w') as f:
        yaml.dump(interconnects, f, default_flow_style=False, sort_keys=False)

    return interconnects

if __name__ == "__main__":
    interconnects_file = Path(__file__).parent.parent.parent / "references" / "interconnects_sample.xlsx"
    output_yml = Path(__file__).parent / "interconnects.yml"
    interconnects_to_yml(interconnects_file, output_yml)
    print(f"Wrote {output_yml}")