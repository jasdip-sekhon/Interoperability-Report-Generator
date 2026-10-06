"""Turn a folder of rendered charts into a slide deck.

Reads the manifest the plotting runner writes rather than the file names, so this script
knows nothing about switches, metrics or test types. A chart type added upstream appears
in the deck with no change here; a capture missing one produces a shorter deck, not a
broken one.

Charts appear in manifest order, grouped by scope - the charts covering the whole capture
first, then one group per switch.
"""

from pathlib import Path
import argparse

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5

# borrowed from the hand-built decks so the generated ones sit at the same measurements
HEADING = (0.35, 0.20, 12.6, 0.42)
RULE = (0.20, 0.82, 12.9, 0.03)
FOOTNOTE = (0.45, 6.35, 12.3, 0.35)

# a cell is (x, caption y, picture y, width, height); a picture is fitted inside it.
# One chart gets the whole slide rather than a lonely quarter of it.
GRID_CELLS = ((0.35, 0.97, 1.24, 5.9, 2.25), (6.85, 0.97, 1.24, 5.9, 2.25),
              (0.35, 3.71, 3.98, 5.9, 2.25), (6.85, 3.71, 3.98, 5.9, 2.25))
SINGLE_CELL = ((0.35, 0.97, 1.30, 12.6, 4.90),)
PER_SLIDE = len(GRID_CELLS)

RULE_COLOUR = RGBColor(0x7F, 0x7F, 0x7F)


def load_manifest(path):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def _blank_slide(presentation):
    # layout 6 is Blank in the default template: no placeholders to work around
    return presentation.slides.add_slide(presentation.slide_layouts[6])


def _add_text(slide, box, text, size, bold=False):
    left, top, width, height = box
    shape = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    return shape


def _add_rule(slide):
    left, top, width, height = RULE
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RULE_COLOUR
    shape.line.fill.background()
    shape.shadow.inherit = False


def _fit(image_path, max_width, max_height):
    """Scale an image to fit a cell without distorting it.

    Chart aspect ratios differ a lot - a 35-entry box plot is wide and short, a
    temperature chart is nearly square - so a fixed size would squash one or the other.
    """
    with Image.open(image_path) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


def _place_chart(slide, chart, cell, charts_dir):
    x, caption_y, picture_y, cell_w, cell_h = cell
    _add_text(slide, (x, caption_y, cell_w, 0.24), chart["title"], 10)
    path = charts_dir / chart["file"]
    width, height = _fit(path, cell_w, cell_h)
    # centre it horizontally in the cell, so narrow charts are not left-stranded
    left = x + (cell_w - width) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(picture_y),
                             Inches(width), Inches(height))


def add_title_slide(presentation, report, manifests):
    slide = _blank_slide(presentation)
    _add_text(slide, (0.9, 2.6, 11.5, 0.9), f"Interoperability Test Report — {report}",
              34, bold=True)
    names = []
    for manifest in manifests:
        names.append(manifest["test_type"])
    _add_text(slide, (0.9, 3.6, 11.5, 0.5), " · ".join(names) or "no testcases", 14)
    return slide


def add_section_slide(presentation, text):
    slide = _blank_slide(presentation)
    _add_text(slide, (0.55, 2.85, 12.2, 0.7), text, 26, bold=True)
    _add_rule(slide)
    return slide


def add_chart_slide(presentation, heading, charts, charts_dir):
    slide = _blank_slide(presentation)
    _add_text(slide, HEADING, heading, 16, bold=True)
    _add_rule(slide)
    if len(charts) == 1:
        cells = SINGLE_CELL
    else:
        cells = GRID_CELLS
    for chart, cell in zip(charts, cells):
        _place_chart(slide, chart, cell, charts_dir)
    _add_text(slide, FOOTNOTE,
              "Plots regenerated from the raw capture; gaps are entries with no "
              "measurement rather than zeroes.", 9)
    return slide


def group_by_scope(charts):
    """[(scope, charts)] in manifest order - the runner decides what order means."""
    groups = []
    for chart in charts:
        if not groups or groups[-1][0] != chart["scope"]:
            groups.append((chart["scope"], []))
        groups[-1][1].append(chart)
    return groups


def paginate(items, size):
    pages = []
    for start in range(0, len(items), size):
        pages.append(items[start:start + size])
    return pages


def _heading_for(scope, page_number, page_count):
    if scope == "all":
        heading = "Across all switches"
    else:
        heading = f"Switch {scope}"
    if page_count > 1:
        heading = f"{heading} (page {page_number} of {page_count})"
    return heading


def add_testcase(presentation, manifest, charts_dir):
    """A section divider and the chart slides under it. Returns the slide count."""
    add_section_slide(presentation, f"{manifest['test_type'].title()} Plots")
    slides = 1
    for scope, charts in group_by_scope(manifest["charts"]):
        pages = paginate(charts, PER_SLIDE)
        for number, page in enumerate(pages, start=1):
            add_chart_slide(presentation, _heading_for(scope, number, len(pages)),
                            page, charts_dir)
            slides = slides + 1
    return slides


def build_deck(manifest_paths, out_path, report="Report"):
    """One deck from one or more testcases, each becoming its own section.

    Charts live next to their own manifest, so testcases rendered into separate folders
    stay separate - which is what lets several appear in one deck without overwriting
    each other's files.
    """
    if isinstance(manifest_paths, (str, Path)):
        manifest_paths = [manifest_paths]

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH)
    presentation.slide_height = Inches(SLIDE_HEIGHT)

    manifests = []
    for path in manifest_paths:
        manifests.append(load_manifest(path))

    add_title_slide(presentation, report, manifests)
    slides = 1
    charts = 0
    for path, manifest in zip(manifest_paths, manifests):
        slides = slides + add_testcase(presentation, manifest, Path(path).parent)
        charts = charts + len(manifest["charts"])

    presentation.save(out_path)
    print(f"Wrote {out_path} - {slides} slides from {charts} chart(s) "
          f"across {len(manifests)} testcase(s)")
    return presentation


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("manifests", nargs="+",
                        help="manifest.yml files, one per testcase, in deck order")
    parser.add_argument("-o", "--out", required=True, help="the .pptx to write")
    parser.add_argument("--report", default="Report", help="name for the title slide")
    args = parser.parse_args()
    build_deck(args.manifests, args.out, args.report)
