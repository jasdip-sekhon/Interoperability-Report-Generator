from pathlib import Path
from openpyxl import load_workbook
import yaml

LOOKUP_SHEET = "Chart label lookup"

def _lookup_sheet(wb):
    # a real topology workbook carries cables, lane links and a link matrix alongside
    # the lookup, and the active sheet is whichever was open when it was saved - so the
    # sheet has to be named, not guessed
    if LOOKUP_SHEET in wb.sheetnames:
        return wb[LOOKUP_SHEET]
    raise ValueError(f"no '{LOOKUP_SHEET}' sheet in this workbook; "
                     f"found {wb.sheetnames}")

def interconnects_to_yml(interconnects_file, output_yml):
    wb = load_workbook(interconnects_file, read_only=True)
    worksheet = _lookup_sheet(wb)

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return {}

    headers = rows[0]
    data_rows = rows[1:]

    interconnects = []
    for row in data_rows:
        if not any(row):
            continue
        record = {}
        for i in range(len(headers)): 
            header = headers[i]
            if i < len(row):
                record[header] = row[i]
        interconnects.append(record)

    # one shared yaml serves whichever report ran last, so stamp which workbook it
    # came from: charting a rack against another rack's axis is otherwise only
    # visible as every interconnect going unmatched
    with open(output_yml, 'w') as f:
        f.write("# source: " + str(Path(interconnects_file).resolve()) + "\n")
        yaml.dump(interconnects, f, default_flow_style=False, sort_keys=False)

    return interconnects

if __name__ == "__main__":
    interconnects_file = Path(__file__).parent.parent / "references" / "interconnects_sample.xlsx"
    output_yml = Path(__file__).parent / "interconnects.yml"
    interconnects_to_yml(interconnects_file, output_yml)
    print(f"Wrote {output_yml}")