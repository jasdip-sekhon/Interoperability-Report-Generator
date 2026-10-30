"""What one report covers, from a config file or from the data itself.

A .yml is read as written. Anything else - one capture, or a folder of them - becomes a
config naming each capture as its own testcase. Nothing is paired: the real captures
spell one testcase 'lane1-1' and 'ln1-1', and a wrong pairing is silent.
"""

from pathlib import Path

import yaml
from openpyxl import load_workbook

# the sheets the converter reads. A topology workbook has none of them, and a folder of
# captures usually has the topology sitting right next to them.
CAPTURE_SHEETS = ("log", "get_link_flap", "get_module_info", "get_switch_sensor",
                  "get_fec_counter")


def load_config(path):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def is_capture(path):
    """Whether this workbook holds measurements, by the sheets it carries."""
    workbook = load_workbook(path, read_only=True)
    names = workbook.sheetnames
    workbook.close()
    for name in names:
        for sheet in CAPTURE_SHEETS:
            if name.endswith(sheet):
                return True
    return False


def _captures_in(folder):
    captures = []
    skipped = []
    for path in sorted(folder.glob("*.xlsx")):
        if is_capture(path):
            captures.append(path)
        else:
            skipped.append(path.name)
    if skipped:
        print(f"== config == skipping {', '.join(skipped)}: no measurement sheets, "
              f"so not a capture")
    if not captures:
        raise SystemExit(f"no captures in {folder}")
    return captures


def config_for(target):
    """A config from whatever the user pointed at: a .yml, a capture, or a folder."""
    target = Path(target)
    if target.suffix.lower() in (".yml", ".yaml"):
        return load_config(target)

    if target.is_dir():
        captures = _captures_in(target)
        report = target.name
    elif target.exists():
        if not is_capture(target):
            raise SystemExit(f"{target.name} has no measurement sheets; it is not a "
                             f"capture. Pass a topology with --topology.")
        captures = [target]
        report = target.stem
    else:
        raise SystemExit(f"no such file or folder: {target}")

    testcases = []
    for capture in captures:
        testcases.append({"name": capture.stem,
                          "captures": [str(capture.resolve())]})
    print(f"== config == none given; charting {len(testcases)} capture(s) from "
          f"{target}, each on its own")
    return {"report": report, "testcases": testcases}
